import asyncio
import logging
from datetime import datetime, time, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from aiogram import Bot, Dispatcher
from aiogram.fsm.storage.base import StorageKey
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

from keyboards import main_keyboard
from sqlwork.orders import delete_pin_orders, get_pin_order, get_pin_orders_for_archive


logger = logging.getLogger(__name__)
KYIV = ZoneInfo("Europe/Kyiv")
ARCHIVE_DIR = Path(__file__).resolve().parents[1] / "payment_history"
RECEIPTS_DIR = ARCHIVE_DIR / "receipts"
WORKBOOK_PATH = ARCHIVE_DIR / "payments.xlsx"
HEADERS = [
    "ID заказа",
    "Дата создания (Киев)",
    "Дата оплаты / подтверждения (Киев)",
    "Telegram ID",
    "Имя пользователя",
    "Username",
    "Закреп, часов",
    "Сумма, грн",
    "Сумма, Stars",
    "Способ оплаты",
    "Статус заказа",
    "Чек / скриншот победы",
    "Telegram charge ID",
]


def seconds_until_kyiv_midnight(now: datetime | None = None) -> float:
    current = now.astimezone(KYIV) if now is not None else datetime.now(KYIV)
    next_date = current.date() + timedelta(days=1)
    next_midnight = datetime.combine(next_date, time.min, tzinfo=KYIV)
    return max(0.0, (next_midnight - current).total_seconds())


def _format_timestamp(timestamp: int | None) -> str:
    if timestamp is None:
        return ""
    return datetime.fromtimestamp(timestamp, KYIV).strftime("%Y-%m-%d %H:%M:%S")


async def _save_receipt(bot: Bot, order: dict) -> tuple[str, str | None]:
    file_id = order.get("payment_proof_file_id")
    if not file_id:
        if order.get("telegram_payment_charge_id"):
            return "Оплата Telegram Stars", None
        return "—", None

    try:
        telegram_file = await bot.get_file(file_id)
        extension = Path(telegram_file.file_path or "").suffix.lower()
        if extension not in {".jpg", ".jpeg", ".png", ".webp", ".pdf"}:
            extension = ".file"

        RECEIPTS_DIR.mkdir(parents=True, exist_ok=True)
        receipt_path = RECEIPTS_DIR / f"order_{order['id']}{extension}"
        await bot.download_file(telegram_file.file_path, destination=receipt_path)
        proof_label = (
            "Открыть скриншот победы"
            if order.get("payment_method") == "giveaway"
            else "Открыть чек"
        )
        return proof_label, receipt_path.relative_to(ARCHIVE_DIR).as_posix()
    except Exception:
        logger.exception("Could not download receipt for order %s", order["id"])
        return f"Telegram file ID: {file_id}", None


def _open_workbook():
    ARCHIVE_DIR.mkdir(parents=True, exist_ok=True)
    if WORKBOOK_PATH.exists():
        workbook = load_workbook(WORKBOOK_PATH)
        sheet = workbook.active
        if [sheet.cell(1, column).value for column in range(1, len(HEADERS) + 1)] != HEADERS:
            raise ValueError("Payment archive has an unexpected column layout")
        return workbook, sheet

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Платежи"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    return workbook, sheet


async def archive_payment_orders(bot: Bot, orders: list[dict]) -> None:
    workbook, sheet = _open_workbook()
    existing_ids = {
        str(sheet.cell(row, 1).value)
        for row in range(2, sheet.max_row + 1)
        if sheet.cell(row, 1).value is not None
    }

    for order in orders:
        order_id = order["id"]
        if str(order_id) in existing_ids:
            continue

        receipt_text, receipt_link = await _save_receipt(bot, order)
        payment_method = {
            "card": "Карта / IBAN",
            "stars": "Telegram Stars",
            "giveaway": "Победитель розыгрыша",
        }.get(order.get("payment_method"), order.get("payment_method", "Неизвестно"))
        row = [
            order_id,
            _format_timestamp(order.get("created_at")),
            _format_timestamp(order.get("payment_submitted_at")),
            order.get("user_id"),
            order.get("user_name") or "",
            f"@{order['username']}" if order.get("username") else "",
            order.get("hours"),
            order.get("amount"),
            order.get("stars_amount"),
            payment_method,
            order.get("status"),
            receipt_text,
            order.get("telegram_payment_charge_id") or "",
        ]
        sheet.append(row)
        if receipt_link:
            receipt_cell = sheet.cell(sheet.max_row, 12)
            receipt_cell.hyperlink = receipt_link
            receipt_cell.font = Font(color="0563C1", underline="single")
        existing_ids.add(str(order_id))

    sheet.auto_filter.ref = sheet.dimensions
    widths = [12, 22, 24, 16, 24, 24, 16, 14, 16, 22, 24, 32, 38]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width

    if sheet.max_row >= 2:
        table_ref = f"A1:M{sheet.max_row}"
        if "PaymentOrders" in sheet.tables:
            sheet.tables["PaymentOrders"].ref = table_ref
        else:
            table = Table(displayName="PaymentOrders", ref=table_ref)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            sheet.add_table(table)

    temporary_path = WORKBOOK_PATH.with_suffix(".tmp.xlsx")
    workbook.save(temporary_path)
    temporary_path.replace(WORKBOOK_PATH)


async def _archive_and_clear(bot: Bot, dispatcher: Dispatcher) -> None:
    orders = get_pin_orders_for_archive()
    if not orders:
        return

    await archive_payment_orders(bot, orders)
    order_ids = [order["id"] for order in orders]
    deleted_count = delete_pin_orders(order_ids)
    if deleted_count != len(order_ids):
        logger.warning(
            "Archived %s payment orders but deleted %s database rows",
            len(order_ids),
            deleted_count,
        )

    from payment_init import stop_payment_countdown

    notified_users: set[int] = set()
    for order in orders:
        order_id = order["id"]
        user_id = order["user_id"]
        if get_pin_order(order_id) is not None:
            continue

        stop_payment_countdown(order_id)
        key = StorageKey(bot_id=bot.id, chat_id=user_id, user_id=user_id)
        state_data = await dispatcher.storage.get_data(key)
        if state_data.get("order_id") != order_id:
            continue

        await dispatcher.storage.set_state(key, None)
        await dispatcher.storage.set_data(key, {})
        if user_id in notified_users:
            continue
        notified_users.add(user_id)
        try:
            await bot.send_message(
                chat_id=user_id,
                text=(
                    "Ваш платёжный заказ перенесён в архив, активные заказы "
                    "очищены. Чтобы оформить новый пост, начните заново."
                ),
                reply_markup=main_keyboard(),
            )
        except Exception:
            logger.exception("Could not notify user %s about order reset", user_id)

    logger.info("Archived and cleared %s payment orders", deleted_count)


async def run_daily_payment_archive(bot: Bot, dispatcher: Dispatcher) -> None:
    while True:
        await asyncio.sleep(seconds_until_kyiv_midnight())
        while True:
            try:
                await _archive_and_clear(bot, dispatcher)
                break
            except asyncio.CancelledError:
                raise
            except Exception:
                logger.exception("Daily payment archive failed; retrying in 10 minutes")
                await asyncio.sleep(600)